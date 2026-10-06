"""Shared helpers: config loading, date formatting, output paths, deterministic file helpers."""
from __future__ import annotations

import datetime as dt
import json
import os
import random
import re
import zipfile
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[2]
SYN = "SYNTHETIC – DEMO ONLY"
SEED = 2026

random.seed(SEED)
np.random.seed(SEED)


def _load(name: str) -> dict:
    with open(ROOT / "config" / name, encoding="utf-8") as fh:
        return json.load(fh)


PERSONAS = _load("personas.json")
TIMELINE = _load("timeline.json")
ROLES = _load("roles.json")

SUBJECT = PERSONAS["subject"]
DECOY = PERSONAS["decoy"]
STUB = PERSONAS["duplicate_stub"]
PEOPLE = PERSONAS["people"]
JOINERS = PERSONAS["other_joiners"]
EV = TIMELINE["events"]
AS_OF = TIMELINE["as_of"]
START = TIMELINE["start_date"]


def out_root() -> Path:
    return Path(os.environ.get("GDTP_OUT", str(ROOT)))


def corpus_path(name: str) -> Path:
    p = out_root() / "corpus"
    p.mkdir(parents=True, exist_ok=True)
    return p / name


def eval_path(name: str) -> Path:
    p = out_root() / "evaluation"
    p.mkdir(parents=True, exist_ok=True)
    return p / name


def meta_path(name: str) -> Path:
    p = out_root() / "metadata"
    p.mkdir(parents=True, exist_ok=True)
    return p / name


def d(iso: str) -> dt.date:
    return dt.date.fromisoformat(iso)


def short(iso: str) -> str:
    """'2026-10-16' -> '16 Oct 2026'"""
    x = d(iso)
    return f"{x.day} {x.strftime('%b')} {x.year}"


def long(iso: str) -> str:
    """'2026-11-02' -> 'Monday 2 November 2026'"""
    x = d(iso)
    return f"{x.strftime('%A')} {x.day} {x.strftime('%B')} {x.year}"


def day_month(iso: str) -> str:
    x = d(iso)
    return f"{x.day} {x.strftime('%b')}"


FIXED_ZIP_TIME = (2026, 10, 27, 9, 0, 0)
FIXED_DATETIME = dt.datetime(2026, 10, 27, 9, 0, 0)


def normalise_zip(path: Path) -> None:
    """Rewrite an OOXML zip with fixed member timestamps and sorted order for reproducible bytes."""
    path = Path(path)
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    items.sort(key=lambda t: (t[0] != "[Content_Types].xml", t[0]))
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            zi = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-10-27T09:00:00Z", data)
            zout.writestr(zi, data)
    os.replace(tmp, path)


# ---------------------------------------------------------------- PDF helpers
def pdf_doc(path: Path, title: str, author: str, subject: str = ""):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate

    return SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=20 * 2.835, rightMargin=20 * 2.835,
        topMargin=22 * 2.835, bottomMargin=22 * 2.835, title=title, author=author,
        subject=subject, invariant=1,
    )


def footer_fn(left_text: str):
    def _f(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillGray(0.35)
        canvas.drawString(20 * 2.835, 12 * 2.835, left_text)
        canvas.drawCentredString(105 * 2.835, 12 * 2.835, SYN)
        canvas.drawRightString(190 * 2.835, 12 * 2.835, f"Page {doc.page}")
        canvas.restoreState()
    return _f


def styles():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("Cell", parent=ss["BodyText"], fontSize=8.5, leading=10.5))
    ss.add(ParagraphStyle("CellB", parent=ss["Cell"], fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("Small", parent=ss["BodyText"], fontSize=8.5, leading=11,
                          textColor="#444444"))
    return ss


def table(rows, col_widths, header=True, style_cmds=None):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    ss = styles()
    data = []
    for r_i, row in enumerate(rows):
        out = []
        for cell in row:
            st = ss["CellB"] if (header and r_i == 0) else ss["Cell"]
            out.append(Paragraph(str(cell), st))
        data.append(out)
    t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    cmds = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if header:
        cmds.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6e9ee")))
    t.setStyle(TableStyle(cmds + (style_cmds or [])))
    return t


# ---------------------------------------------------------------- Excel helpers
def new_workbook():
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.creator = "GDTP synthetic data generator"
    wb.properties.lastModifiedBy = "GDTP synthetic data generator"
    wb.properties.created = FIXED_DATETIME
    wb.properties.modified = FIXED_DATETIME
    return wb


def style_header(ws, ncols: int, widths=None):
    from openpyxl.styles import Font, PatternFill

    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE3EC")
    ws.freeze_panes = "A2"
    for i in range(ncols):
        w = (widths or [])[i] if widths and i < len(widths) else 18
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
    ws.oddFooter.center.text = SYN
    ws.oddHeader.center.text = SYN


def save_workbook(wb, path: Path):
    wb.save(path)
    normalise_zip(path)


# ---------------------------------------------------------------- Email helper
def make_eml(path: Path, *, sender, to, cc=None, subject, date, body, message_id, extra=None):
    from email.message import EmailMessage
    from email import policy

    m = EmailMessage(policy=policy.SMTP)
    m["From"] = sender
    m["To"] = to
    if cc:
        m["Cc"] = cc
    m["Subject"] = subject
    m["Date"] = date
    m["Message-ID"] = message_id
    m["X-Synthetic"] = "demo-only"
    for k, v in (extra or {}).items():
        m[k] = v
    m.set_content(body, cte="quoted-printable")
    with open(path, "wb") as fh:
        fh.write(m.as_bytes())
